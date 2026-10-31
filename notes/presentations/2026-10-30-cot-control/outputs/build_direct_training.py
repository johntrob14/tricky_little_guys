from pathlib import Path
import json, html, shutil
import matplotlib
matplotlib.use('Agg')
import matplotlib.pyplot as plt
import numpy as np
from pptx import Presentation
from pptx.util import Inches,Pt
from pptx.dml.color import RGBColor
O=Path('outputs');data=json.loads((O/'direct_training_data.json').read_text());d=data['main'];b=data['backup']
INK='202A36';BLUE='3467A4';TEAL='008F85';ORANGE='D67B35'
plt.rcParams.update({'font.family':'DejaVu Sans','font.size':19,'svg.fonttype':'none','axes.spines.top':False,'axes.spines.right':False,'axes.spines.left':False,'text.color':'#'+INK,'axes.labelcolor':'#'+INK,'xtick.color':'#'+INK,'ytick.color':'#'+INK})
fig,ax=plt.subplots(figsize=(16,4));fig.subplots_adjust(left=.10,right=.985,bottom=.24,top=.97)
ms=['base','grpo3','lang1-s40','len2'];v=np.array([d[m]['mean_strict']*100 for m in ms]);ci=np.array([d[m]['ci95'] for m in ms])*100
ax.bar(range(4),v,width=.57,color=['#8995A3','#'+BLUE,'#'+TEAL,'#'+ORANGE],zorder=3)
ax.errorbar(range(4),v,yerr=np.array([v-ci[:,0],ci[:,1]-v]),fmt='none',ecolor='#202A36',elinewidth=1.8,capsize=6,zorder=4)
for i in range(4):ax.text(i,ci[i,1]+2,f'{v[i]:.1f}%',ha='center',va='bottom',fontsize=21)
ax.set_xticks(range(4),['Base','Six CoT rules\nGRPO','Reasoning language\nGRPO · step 40','CoT token budget\nGRPO']);ax.set_ylabel('Strict CoT compliance');ax.set_ylim(0,85);ax.set_yticks([0,20,40,60,80],['0%','20%','40%','60%','80%']);ax.grid(axis='y',color='#E5E8EC',zorder=0);ax.tick_params(length=0,pad=10)
fig.savefig(O/'direct_training.svg',transparent=True);fig.savefig(O/'direct_training.png',dpi=160,facecolor='white');plt.close(fig)
cap=('GPT-OSS-120B; same 480 question–rule items across eight deterministic rule types, with meta discussion allowed. '
     'Bars average samples within each item, then weight items equally. Error bars: 95% question-cluster bootstrap confidence intervals '
     '(413 question clusters, 5,000 draws, seed 20261030); evaluation uncertainty only, one training run per recipe. '
     'Base/grpo3/len2: 1,536 generations each; lang1-s40: 960. Checkpoints: grpo3 s100, lang1 s40, len2 s60. '
     'The metric does not require a correct answer; excluding incomplete generations leaves these four rates unchanged under the saved completion checks.')
slides=[dict(title='Learning to Control Thoughts Directly',bullets=[
 'We apply GRPO directly to the CoT: six lexical/formatting rules, a requested language, or a token budget.',
 'Language and token-budget objectives also improve broader CoT instruction following, suggesting that control can develop under ordinary reasoning training.'
],notes=(
 'We trained GPT-OSS-120B with reasoning enabled using three separate GRPO recipes starting from the base model. '
 'The six-rule recipe rewards lexical and formatting constraints inside the CoT, weights that reward by answer correctness, and penalizes deviations from base reasoning length. '
 'The language recipe rewards French, Spanish, or German reasoning with correctness and length terms, while the token-budget recipe rewards closeness to a requested reasoning length with correctness weighting. '
 'On the same 480 CoT-control items, strict compliance is 3.5% for base, 66.9% after six-rule training, 26.5% after language training, and 12.4% after token-budget training. '
 'These developed recipes demonstrate substantial changes in surface control, while single training runs, shorter reasoning in some arms, and lower constrained answer accuracy limit stronger causal or capability-preservation claims.'
),caption=cap,figure=True),
dict(title='Backup: rewrite SFT teaches CoT constraints',bullets=[
 'Rewrite the model’s own reasoning into requested styles, then train on those traces and final answers alongside plain examples.',
 'This provides explicit examples of compliant reasoning; correct-and-compliant counts help distinguish useful control from surface success.'
],notes=(
 'For Inkling-Small, we fine-tuned the base model for two epochs on 384 examples containing rewritten or transformed reasoning and plain traces. '
 'The completion-only supervised loss trains both the native thinking message and the final answer. '
 'At matched thinking effort 0.6 on 96 deterministic Astra items, strict CoT compliance rises from 3/96 to 18/96. '
 'Correct, valid, and compliant responses rise from 2/96 to 8/96, while answer correctness falls from 61/96 to 53/96. '
 'This is a single SFT run with shorter constrained traces and no matched control-data fine-tune, so it establishes a recipe result without isolating which part of the training data caused it.'
),headers=['Inkling-Small','Strict CoT passes','Correct + valid + compliant','Correct answers'],rows=[['Base','3/96','2/96','61/96'],['Rewrite SFT, epoch 2','18/96','8/96','53/96']],caption='One generation per item; original Astra deterministic subset; effort 0.6, temperature 1.0. Strict rule scoring allows meta discussion. Raw regrade and independent count aggregation; no new semantic-judge validation.'),
dict(title='Backup: self-distillation, then direct RL',bullets=[
 'OPSD trains a plain-prompted Qwen student to match the same base model given a constraint-specific self-instruction prefill.',
 'Subsequent GRPO adds explicit CoT-compliance reward; the joint metric checks whether compliant traces also yield valid, correct answers.'
],notes=(
 'For Qwen3.8-27B, on-policy self-distillation scores student-generated continuations under the same base model given a constraint-specific self-instruction prefill. '
 'The reverse-KL objective trains the sampled continuation, including reasoning, without using a compliance reward during distillation. '
 'We then apply GRPO to the selected OPSD checkpoint using CoT-rule rewards, correctness weighting, and a base-relative length penalty. '
 'On 96 deterministic Astra items, raw strict compliance rises from 0/96 to 20/96 after OPSD and 51/96 after GRPO, while correct, valid, and compliant responses rise from 0/96 to 6/96 to 26/96. '
 'Four OPSD rule passes lack valid completions, sampled passes include degenerate traces, and the selected single-run checkpoints have shorter reasoning and lower answer accuracy than base.'
),headers=['Qwen3.8-27B','Strict CoT passes','Correct + valid + compliant','Correct answers'],rows=[['Base','0/96','0/96','57/96'],['OPSD, step 30','20/96','6/96','48/96'],['OPSD + GRPO, 30 RL steps','51/96','26/96','50/96']],caption='One generation per item; original Astra deterministic subset; temperature 1.0, top-p 0.95, 25k cap, thinking wording, no evaluation prefill. OPSD s30 was selected after evaluating multiple checkpoints. Four OPSD raw passes fail completion checks, leaving 16/96 completed strict passes.')]
prs=Presentation();prs.slide_width=Inches(13.333);prs.slide_height=Inches(7.5)
def txt(s,t,x,y,w,h,size=23,bold=False,color=INK):
 sh=s.shapes.add_textbox(Inches(x),Inches(y),Inches(w),Inches(h));tf=sh.text_frame;tf.word_wrap=True;tf.margin_left=tf.margin_right=Inches(.02)
 p=tf.paragraphs[0];p.text=t;p.font.name='Aptos';p.font.size=Pt(size);p.font.bold=bold;p.font.color.rgb=RGBColor.from_string(color)
 return sh
for i,z in enumerate(slides):
 s=prs.slides.add_slide(prs.slide_layouts[6]);txt(s,z['title'],.65,.35,12,1.05,32,True)
 if z.get('figure'):
  s.shapes.add_picture(str(O/'direct_training.png'),Inches(.65),Inches(1.55),width=Inches(12))
  txt(s,'GPT-OSS-120B · 480 items · 95% CIs · one run per recipe',.65,1.1,12,.35,15,color='65717E')
  for j,t in enumerate(z['bullets']):txt(s,'• '+t,.65,4.85+j*.82,12,.75,23)
  txt(s,'Strict rule compliance allows meta discussion; answer correctness is measured separately.',.65,6.95,12,.3,12,color='65717E')
 else:
  for j,t in enumerate(z['bullets']):txt(s,'• '+t,.65,1.5+j*1.05,12,.95,23)
  table=s.shapes.add_table(len(z['rows'])+1,4,Inches(.65),Inches(3.85),Inches(12),Inches(2.0)).table
  for j,w in enumerate([3.7,2.4,3.5,2.4]):table.columns[j].width=Inches(w)
  for ri,row in enumerate([z['headers']]+z['rows']):
   for ci,vv in enumerate(row):
    cell=table.cell(ri,ci);cell.text=vv;cell.fill.solid();cell.fill.fore_color.rgb=RGBColor.from_string('E8F1F5' if ri==0 else 'F4F6F8');cell.margin_top=Inches(.08)
    for p in cell.text_frame.paragraphs:p.font.name='Aptos';p.font.size=Pt(17);p.font.bold=ri==0;p.font.color.rgb=RGBColor.from_string(INK)
  txt(s,z['caption'],.65,6.15,12,.9,14,color='65717E')
 s.notes_slide.notes_text_frame.text=z['notes']+'\n\n'+z['caption']+'\n\nSources: direct_training_data.json and direct_training_recount.py. Run recipes: rl/grpo.py, rl/grpo_types.py, rl/grpo_length.py, rl/opsd.py, rl/inksm_sft.py in the recorded source checkouts.'
prs.save(O/'learning_to_control_thoughts.pptx')
(O/'learning_to_control_thoughts.json').write_text(json.dumps(slides,indent=2))
(O/'learning_to_control_thoughts.txt').write_text('\n\n'.join(z['title']+'\n'+'\n'.join('• '+t for t in z['bullets'])+'\n\nDescription / speaker notes\n'+z['notes']+'\n\nCaption / evidence details\n'+z['caption']+ ('\n\n'+'\n'.join(' | '.join(r) for r in [z['headers']]+z['rows']) if 'rows' in z else '') for z in slides))
sections=[]
for z in slides:
 viz=(O/'direct_training.svg').read_text() if z.get('figure') else '<table>'+''.join('<tr>'+''.join('<td>'+html.escape(v)+'</td>' for v in row)+'</tr>' for row in [z['headers']]+z['rows'])+'</table>'
 sections.append('<section><h1>'+html.escape(z['title'])+'</h1>'+viz+'<ul>'+''.join('<li>'+html.escape(t)+'</li>' for t in z['bullets'])+'</ul><p><b>Description / notes</b><br>'+html.escape(z['notes'])+'</p><small>'+html.escape(z['caption'])+'</small></section>')
(O/'learning_to_control_thoughts.html').write_text('<!doctype html><html><head><meta charset="utf-8"><style>body{font:21px/1.4 Arial;color:#202a36;background:#e9edf1;margin:0}section{max-width:1120px;background:white;padding:40px;margin:25px auto}h1{font-size:34px}svg{width:100%;height:auto}li{margin:15px 0}p{font-size:18px}small{font-size:14px}table{border-collapse:collapse;width:100%;font-size:20px}td{padding:18px;border-bottom:1px solid #ddd}tr:first-child{font-weight:bold;background:#e8f1f5}</style></head><body>'+''.join(sections)+'</body></html>')
for filename in ['build_direct_training.py','direct_training_recount.py']:shutil.copy(Path('work')/filename,O/filename)
r=Presentation(O/'learning_to_control_thoughts.pptx');assert len(r.slides)==3
for s in r.slides:
 for sh in s.shapes:assert sh.left>=0 and sh.top>=0 and sh.left+sh.width<=r.slide_width and sh.top+sh.height<=r.slide_height
print('Created one main slide and two backups, SVG/PNG, descriptions and notes, data, and reproducible scripts.')
