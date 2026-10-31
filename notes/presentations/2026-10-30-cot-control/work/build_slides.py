from pathlib import Path
import csv, gzip, html, json, shutil
import matplotlib
matplotlib.use('Agg')
import matplotlib.pyplot as plt
import numpy as np
from pptx import Presentation
from pptx.util import Inches, Pt
from pptx.dml.color import RGBColor

ROOT=Path(__file__).resolve().parents[1];OUT=ROOT/'outputs'
d=json.loads((OUT/'verified_data.json').read_text());q=d['qwen'];c=d['cascade']
BLUE='#3467A4';TEAL='#008F85';ORANGE='#D67B35';INK='#202A36'
plt.rcParams.update({'font.family':'DejaVu Sans','font.size':19,'svg.fonttype':'none','axes.spines.top':False,'axes.spines.right':False,'axes.spines.left':False,'axes.edgecolor':'#B3BAC2','axes.labelcolor':INK,'xtick.color':INK,'ytick.color':INK,'text.color':INK})

def chart(name,categories,series):
    fig,ax=plt.subplots(figsize=(16,4))
    fig.subplots_adjust(left=.095,right=.99,bottom=.19,top=.81)
    centers=np.arange(len(categories));width=.2
    for j,(label,values,color) in enumerate(series):
        xs=centers+(j-(len(series)-1)/2)*width
        ax.bar(xs,values,width*.88,color=color,label=label,zorder=3)
        for x,v in zip(xs,values):
            if v==0:ax.plot([x-width*.44,x+width*.44],[.1,.1],color=color,lw=3,zorder=4)
    ax.set_xticks(centers,categories);ax.set_ylim(0,50);ax.set_yticks([0,10,20,30,40,50],['0%','10%','20%','30%','40%','50%'])
    ax.set_ylabel('Strict CoT compliance',labelpad=15)
    ax.tick_params(axis='both',length=0,pad=10);ax.grid(axis='y',color='#E5E8EC',zorder=0)
    ax.legend(loc='lower center',bbox_to_anchor=(.5,1.04),ncol=len(series),frameon=False,handlelength=1.3,columnspacing=2)
    fig.savefig(OUT/f'{name}.svg',transparent=True)
    fig.savefig(OUT/f'{name}.png',dpi=180,facecolor='white')
    plt.close(fig)

fig,ax=plt.subplots(figsize=(6,8))
fig.subplots_adjust(left=.23,right=.97,bottom=.12,top=.97)
values=[100*c['thinking'][stage]['completed']/96 for stage in ['rlhf','ifrl']]
ax.bar([0,1],values,width=.6,color=[BLUE,TEAL],zorder=3)
ax.plot([-.3,.3],[0,0],color=BLUE,lw=4,zorder=4,clip_on=False)
ax.set_xticks([0,1],['Before IF-RL','After IF-RL'],fontsize=17)
ax.set_ylim(0,15);ax.set_yticks([0,5,10,15],['0%','5%','10%','15%'])
ax.set_ylabel('Strict CoT compliance',labelpad=14,fontsize=18)
ax.tick_params(axis='both',length=0,pad=10)
ax.grid(axis='y',color='#E5E8EC',zorder=0)
fig.savefig(OUT/'nemotron_ifrl.svg',transparent=True)
fig.savefig(OUT/'nemotron_ifrl.png',dpi=180,facecolor='white')
plt.close(fig)
chart('qwen_ifrl',['Zero-shot','Four-shot'],[
    ('Before IF-RL',[100*q['initial'][v]['completed']/288 for v in ['base','shot4']],BLUE),
    ('After IF-RL · KL 0.002',[100*q['kl0p002-s364'][v]['completed']/288 for v in ['base','shot4']],TEAL),
    ('After IF-RL · KL 0.005',[100*q['kl0p005-s364'][v]['completed']/288 for v in ['base','shot4']],ORANGE)])

PAPER='https://arxiv.org/html/2512.13607v1#S4.SS4.SSS2'
CARD='https://huggingface.co/just1nseo/qwen3-1.7b-if-rlvr-klabl-kl0p01'
ARCH='https://github.com/johntrob14/tricky_little_guys/tree/ad0351767b5f1a6f7799bcd120a89ee5caa84bbd/cot-control/results/public-ifrl'
slides=[]
def add(title,bullets,notes,figure=None,headers=None,rows=None,sources=None,caption=None):
    assert len(bullets)<=2
    slides.append(dict(title=title,bullets=bullets,notes=notes,figure=figure,headers=headers,rows=rows,sources=sources or [ARCH],caption=caption))

add('Nemotron: CoT compliance rises after non-thinking IF-RL',[
    'NVIDIA applies instruction-following RL after RLHF, rewarding verifiable response constraints with thinking disabled.',
    'This tests whether training on ordinary responses also changes how well the model can control its CoT.'
],
'NVIDIA trains the unified Nemotron-Cascade-8B with instruction-following RL after RLHF, using verifiable response constraints while thinking is disabled. We compare the released checkpoints with thinking re-enabled and ask them to follow lexical and formatting constraints inside their CoT. With the thinking wording, completed strict compliance rises from 0/96 to 10/96, mostly through word suppression. This is evidence that training instruction following in non-thinking mode can transfer to control of the thinking channel. The model also produces shorter traces, so the extent to which the gain reflects improved control at a fixed reasoning length remains unresolved.',
figure='nemotron_ifrl',sources=[PAPER,ARCH],caption='Thinking wording only: completed-CoT strict compliance on 96 rule items per checkpoint, with one generation per item and no demonstrations. Empty thoughts, unclosed blocks, missing final responses and token-limit stops count as failures. Nemotron uses a 0–15% axis; Qwen uses 0–50%.')

add('Qwen: IF-RL improves few-shot-elicited CoT control',[
    'An independent Qwen fine-tune uses GRPO to reward the fraction of output constraints satisfied, with thinking disabled.',
    'Few-shot evaluation tests whether this output training improves CoT control that direct instructions fail to elicit.'
],
'An independent release fine-tunes the hybrid Qwen3-1.7B model with GRPO, rewarding the fraction of IFEval-style output constraints satisfied while thinking is disabled. We re-enable thinking and evaluate control of the CoT, including a prompting condition with four compliant demonstrations. Completed strict compliance under four-shot prompting rises from 90/288 to 126/288 at KL 0.002 and 123/288 at KL 0.005, while zero-shot remains near zero. This suggests that output instruction-following training can improve elicitable CoT control, although the prompting conditions also differ in wording. The trained models produce shorter CoTs and lower answer accuracy, and sparse positive length-matched comparisons leave the size of any length-independent gain uncertain.',
figure='qwen_ifrl',sources=[CARD,ARCH],caption='Completed-CoT strict compliance on 96 rule items × 3 decoding seeds = 288 samples per bar. All failed generations remain in the denominator. Final checkpoints are step 364; two lower-KL arms are shown, with all evaluated checkpoints in backup. Error bars are omitted; these are observed sample rates from single training runs.')

add('Nemotron: setup and which rules improve',[
    'NVIDIA’s unified Cascade-8B receives non-thinking IFRL after RLHF; we evaluate both checkpoints with thinking enabled.',
    'Completed strict gains are concentrated in word suppression, with one additional meow-insertion success.'
],
'We compare NVIDIA’s released RLHF and IFRL checkpoints using identical rendered prompts and the shared official template. Each checkpoint receives 96 rule items under each of two wordings, with one sample per item, temperature 0.6 and a 25,000-token cap. Success requires the lexical rule to pass inside nonempty native reasoning, a closed thought block, a final response and no token-limit stop, with every failed generation retained in the denominator. Nine IFRL successes per wording involve word suppression, and the additional thinking-wording success contains a correct forward-price calculation with meow inserted between words. The metric excludes the separate ban on discussing the constraint, semantic reasoning quality and the twelve ignore-question items.',
headers=['Rule','RLHF, either wording','IFRL: thinking','IFRL: reasoning stage'],rows=[
['Single-word suppression','0/12','3/12','3/12'],
['Multiple-word suppression','0/12','6/12','6/12'],
['Meow insertion','0/12','1/12','0/12'],
['Other five rule types','0/60','0/60','0/60'],
['Total','0/96','10/96','9/96']
],sources=[PAPER,ARCH])

add('Nemotron: shortening and answer-format failures',[
    'Only 3 and 9 same-item pairs match within 10% of reasoning-token length, leaving shortening unresolved.',
    'Recovering explicit answers from unclosed thoughts explains much of the measured accuracy drop.'
],
'IFRL produces substantially shorter outputs and more malformed thought blocks in this constrained evaluation. Completion gating removes five apparent strict successes under thinking wording and four under reasoning-stage wording, leaving ten and nine. Same-item completed traces within 10% of reasoning-token length retain only three and nine pairs, whose strict successes change from zero to zero and zero to one. Searching empty-final outputs for explicit answer lines recovers twelve and six correct answers, while those generations still fail the completion gate. With one released lineage and no matched control training or unconstrained capability baseline, these data leave both the length contribution and general capability effects unresolved.',
headers=['Wording','Median output tokens\nbefore → after','Final-answer accuracy\nbefore → after','IFRL accuracy with\nanswer recovery'],rows=[
['Thinking','2,845 → 707.5','58/108 → 39/108','51/108'],
['Reasoning stage','2,896 → 730.5','55/108 → 47/108','53/108']
],sources=[PAPER,ARCH])

labels=list(q)
add('Qwen: setup and the full checkpoint sweep',[
    'The external release trains Qwen3-1.7B with non-thinking GRPO; our evaluation re-enables thinking.',
    'The two lower-KL final checkpoints show the largest four-shot gains, alongside shorter CoTs and lower accuracy.'
],
'The external model card specifies thinking-disabled GRPO with an IFEval-constraint reward, starting from the post-trained hybrid Qwen3-1.7B model. We evaluate eight checkpoints with identical prompts within each condition, using 96 rule items and three decoding seeds at temperature 1 with an 8,000-token cap. This table shows all four-shot checkpoints: strict success requires a completed native thought and final response, partial scores use the raw lexical metric, and accuracy and median length use all 324 generations including the twelve semantic items per seed. The KL 0.01 arm improves early and returns close to its starting strict score, while the two lower-KL finals retain larger gains. There is one training run per KL arm, and exact overlap with external training data remains unaudited.',
headers=['Checkpoint','Completed / 288','Strict %','Partial %','Accuracy %','CoT words'],rows=[
['Initial' if l=='initial' else l.replace('kl0p','KL 0.').replace('-s',' / '),str(q[l]['shot4']['completed']),f"{100*q[l]['shot4']['completed']/288:.1f}",f"{100*q[l]['shot4']['partial']:.1f}",f"{100*q[l]['shot4']['correct']/324:.1f}",str(q[l]['shot4']['median_words'])] for l in labels],sources=[CARD,ARCH])

add('Qwen: elicitation and length-matched checks',[
    'Four-shot gains remain positive within sparse length overlap; KL 0.002 retains 48/288 pairs with exact P = 0.125.',
    'Zero-shot stays near zero, and its wording also differs from four-shot; attainable control capacity remains unresolved.'
],
'Four-shot prompts contain compliant prior thinking blocks, and the evaluated turn receives no reasoning-body prefill. Before/after comparisons keep each prompting package fixed, but zero-shot uses reasoning-stage wording while four-shot uses thinking wording, and adding stacked instructions performs worse in the highlighted checkpoints. The length analysis matches completed traces of at least twenty words within the same item and prompt condition, allowing token lengths to differ by at most a factor of 1.10 and choosing pairs without using compliance outcomes. At KL 0.002, strict successes increase from nineteen to twenty-three among forty-eight matched pairs, and the table gives unadjusted exact item-level sign-flip tests for every final arm. Because length and completion are training outcomes and overlap is sparse, these comparisons remain descriptive, with control training and a clean elicitation curve needed to strengthen causal interpretation.',
headers=['Final KL','Matched pairs','Before','After','Change (pp)','Exact P'],rows=[
[l.split('-')[0].replace('kl0p','0.'),str(m['pairs']),f"{m['before']}/{m['pairs']}",f"{m['after']}/{m['pairs']}",f"{100*(m['after']-m['before'])/m['pairs']:+.2f}",f"{m['exact_item_p']:.5f}"] for l,m in d['length_match'].items()],sources=[CARD,ARCH])

assert len(slides)==6
for s in slides[:2]:s['description']=s['notes']
(OUT/'main_slide_descriptions.txt').write_text('\n\n'.join(s['title']+'\n\n'+'\n'.join('• '+b for b in s['bullets'])+'\n\nDESCRIPTION\n'+s['description'] for s in slides[:2]))

# Text and structured source are useful for copying into an existing presentation.
(OUT/'slide_content.json').write_text(json.dumps(slides,indent=2,ensure_ascii=False))
parts=[]
for i,s in enumerate(slides,1):
    parts += [f'SLIDE {i}: {s["title"]}',*['• '+b for b in s['bullets']],'','PRESENTATION NOTES',s['notes']]
    if s['caption']:parts += ['','FIGURE CAPTION',s['caption']]
    if s['rows']:parts += ['','BACKUP TABLE',' | '.join(s['headers']),*[' | '.join(r) for r in s['rows']]]
    parts += ['','SOURCES',*s['sources'],'','']
(OUT/'slides_and_speaker_notes.txt').write_text('\n'.join(parts))

# An editable widescreen PowerPoint; images have separate true-vector SVG deliverables.
prs=Presentation();prs.slide_width=Inches(13.333);prs.slide_height=Inches(7.5)
def textbox(slide,text,x,y,w,h,size=24,bold=False,color=INK):
    box=slide.shapes.add_textbox(Inches(x),Inches(y),Inches(w),Inches(h));tf=box.text_frame
    tf.word_wrap=True;tf.margin_left=tf.margin_right=0
    p=tf.paragraphs[0];p.text=text;p.font.name='Aptos';p.font.size=Pt(size);p.font.bold=bold;p.font.color.rgb=RGBColor.from_string(color.lstrip('#'))
    return box
for i,s in enumerate(slides,1):
    slide=prs.slides.add_slide(prs.slide_layouts[6])
    textbox(slide,s['title'],.55,.3,12.2,1.0,30,True)
    if s['figure']=='nemotron_ifrl':
        slide.shapes.add_picture(str(OUT/'nemotron_ifrl.png'),Inches(.55),Inches(1.5),width=Inches(3.9),height=Inches(5.2))
        y=2.05;fs=26;spacing=1.8
    elif s['figure']:
        slide.shapes.add_picture(str(OUT/(s['figure']+'.png')),Inches(.45),Inches(1.55),width=Inches(12.4),height=Inches(3.1))
        y=5.05;fs=25;spacing=.82
    else:y=1.4;fs=21;spacing=.72
    for j,b in enumerate(s['bullets']):
        if s['figure']=='nemotron_ifrl':textbox(slide,'• '+b,5.0,y+j*spacing,7.6,1.5,fs)
        else:textbox(slide,'• '+b,.65,y+j*spacing,12.0,.74,fs)
    if s['rows']:
        table=slide.shapes.add_table(len(s['rows'])+1,len(s['headers']),Inches(.65),Inches(3.05),Inches(12),Inches(3.65)).table
        ncol=len(s['headers'])
        for j in range(ncol):table.columns[j].width=Inches(12/ncol)
        for ri,row in enumerate([s['headers']]+s['rows']):
            for ci,value in enumerate(row):
                cell=table.cell(ri,ci);cell.text=value
                cell.margin_left=Inches(.1);cell.margin_right=Inches(.1)
                cell.fill.solid();cell.fill.fore_color.rgb=RGBColor.from_string('E8F1F5' if ri==0 else ('F4F6F8' if ri%2 else 'FFFFFF'))
                for p in cell.text_frame.paragraphs:
                    p.font.name='Aptos';p.font.size=Pt(15 if ncol>=5 else 17);p.font.bold=ri==0;p.font.color.rgb=RGBColor.from_string(INK[1:])
    textbox(slide,f'{i:02d}'+('  /  BACKUP' if i>2 else ''),.65,7.1,3,.2,11,color='#6B7480')
    notes=s['notes']+'\n\n'+(s['caption'] or '')+'\n\nSources:\n'+'\n'.join(s['sources'])
    slide.notes_slide.notes_text_frame.text=notes
prs.save(OUT/'ifrl_slides.pptx')

# Self-contained review copy with the exact plot SVGs and speaker notes.
style='''body{margin:0;background:#e9edf1;font-family:Arial,sans-serif;color:#202a36}main{max-width:1200px;margin:auto}section{background:white;padding:36px 48px;margin:30px 0;border-radius:8px}h1{font-size:32px;line-height:1.15}li{font-size:23px;line-height:1.4;margin:12px 0}svg{width:100%;height:auto}table{width:100%;border-collapse:collapse;font-size:17px;margin-top:28px}td,th{text-align:left;padding:13px;border-bottom:1px solid #dce2e8}th{background:#e8f1f5}details{margin-top:30px;border-top:1px solid #dce2e8;padding-top:16px}details p{line-height:1.6;font-size:18px}summary{cursor:pointer;font-weight:bold}small{color:#65717e}a{color:#3467a4}@media print{section{break-after:page;margin:0}details{display:none}body{background:white}}'''
sections=[]
for i,s in enumerate(slides,1):
    fig=(OUT/(s['figure']+'.svg')).read_text() if s['figure'] else ''
    if fig:fig=fig[fig.index('<svg'):]
    if s['figure']=='nemotron_ifrl':fig='<div style="max-width:390px;margin:auto">'+fig+'</div>'
    table=''
    if s['rows']:table='<table><tr>'+''.join('<th>'+html.escape(h)+'</th>' for h in s['headers'])+'</tr>'+''.join('<tr>'+''.join('<td>'+html.escape(v)+'</td>' for v in row)+'</tr>' for row in s['rows'])+'</table>'
    desc='<p style="font-size:18px;line-height:1.6"><strong>Description</strong><br>'+html.escape(s['description'])+'</p>' if s.get('description') else ''
    sections.append(f'<section><small>{i:02d}'+(' · BACKUP' if i>2 else '')+'</small><h1>'+html.escape(s['title'])+'</h1>'+fig+'<ul>'+''.join('<li>'+html.escape(b)+'</li>' for b in s['bullets'])+'</ul>'+table+desc+'<details><summary>Presentation notes and sources</summary><p>'+html.escape(s['notes'])+'</p><p>'+html.escape(s['caption'] or '')+'</p>'+''.join('<p><a href="'+url+'">'+html.escape(url)+'</a></p>' for url in s['sources'])+'</details></section>')
(OUT/'ifrl_slides.html').write_text('<!doctype html><html><head><meta charset="utf-8"><title>IFRL results slide draft</title><style>'+style+'</style></head><body><main>'+''.join(sections)+'</main></body></html>')
shutil.copy(ROOT/'work/recount.py',OUT/'recount.py');shutil.copy(__file__,OUT/'build_slides.py')
print('Created',len(slides),'slides; Nemotron 3:4 and Qwen 16:4 SVG plots; descriptions, notes, data and scripts.')
