from pathlib import Path
import html, json, shutil
from pptx import Presentation
from pptx.util import Inches, Pt
from pptx.dml.color import RGBColor
from pptx.enum.shapes import MSO_SHAPE
from pptx.enum.text import PP_ALIGN, MSO_ANCHOR
O=Path('outputs')
title='Feedback spillover may improve CoT control with reasoning on'
bullets=[
 'We reward final-answer correctness and output constraints, but apply the RL advantage to both CoT and response tokens; compliant reasoning could help earn that reward.',
 'GPT-OSS-120B: mean partial score 28.1% → 31.7%; strict compliance 3.5% → 2.5% (480 items, one run). Weak evidence for spillover.'
]
# State the update directly, avoiding a contrast construction.
bullets[0]='With reasoning enabled, we reward final-answer correctness and output constraints and apply the RL advantage to both CoT and response tokens.'
notes=(
 'In outrl1, we trained GPT-OSS-120B with reasoning enabled, grading constraints on the final explanation and applying the resulting RL advantage to all generated tokens. '
 'Because the final response depends on the preceding CoT, reasoning that follows the constraint could make a rewarded response more likely and be reinforced. '
 'Updates to shared parameters could also transfer output instruction following into CoT control. '
 'Our single run shows a small increase in partial CoT compliance without a clear improvement in strict compliance. '
 'This is weak evidence for the hypothesis: partial scores can credit discussion of the rule, and we have no matched control that separates spillover from shortening or other training effects.'
)
source='https://arxiv.org/abs/2511.11584'
audit=(
 'Verification: local recount of 480 deterministic-rule items, averaging samples within each item; semantic ignore_question items excluded. '
 'Partial score: 28.14% to 31.73% (+3.59 percentage points); strict score: 3.52% to 2.50% (-1.02 percentage points). '
 'The partial-score implementation was reused; aggregation was recomputed. These are exploratory scores from one training run, not independent replications. '
 'Forty seeded excerpts reviewed; no full answer-correctness audit. '
 'Training also requires a nonempty reasoning segment and valid completion, so only the constraint grading is final-response-only. '
 'The broader nine-type aggregate in earlier notes uses a different denominator. See spillover_verification.json for this recount.\n'
 'Code: /home/john/code/nocotrl-b-pilot/cot-control/rl/grpo_out.py and rl/grpo.py.\n'
 'Data: /home/john/code/nocotrl-b-pilot/cot-control/results/outctl/{more,big}/cot/{base,outrl1-s53}/graded.json.'
)
prs=Presentation();prs.slide_width=Inches(13.333);prs.slide_height=Inches(7.5)
s=prs.slides.add_slide(prs.slide_layouts[6]);ink='202A36';blue='3467A4';teal='008F85'
def text(t,x,y,w,h,size=24,bold=False,color=ink,align=None):
 sh=s.shapes.add_textbox(Inches(x),Inches(y),Inches(w),Inches(h));tf=sh.text_frame;tf.word_wrap=True
 tf.margin_left=tf.margin_right=Inches(.02)
 p=tf.paragraphs[0];p.text=t;p.font.name='Aptos';p.font.size=Pt(size);p.font.bold=bold;p.font.color.rgb=RGBColor.from_string(color)
 if align is not None:p.alignment=align
 return sh
text(title,.65,.35,12.0,1.25,32,True)
for label,x,w in [('CoT',.85,3.0),('Final response',5.05,3.0),('Reward',9.25,3.0)]:
 sh=s.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE,Inches(x),Inches(1.95),Inches(w),Inches(.85))
 sh.fill.solid();sh.fill.fore_color.rgb=RGBColor.from_string('E8F1F5');sh.line.fill.background()
 tf=sh.text_frame;tf.vertical_anchor=MSO_ANCHOR.MIDDLE
 p=tf.paragraphs[0];p.text=label;p.alignment=PP_ALIGN.CENTER;p.font.name='Aptos';p.font.size=Pt(24);p.font.color.rgb=RGBColor.from_string(ink)
for x in [4.0,8.2]:text('→',x,2.02,.85,.55,28,False,blue,PP_ALIGN.CENTER)
text('Reasoning enabled',.85,2.9,3,.4,16,False,'65717E',PP_ALIGN.CENTER)
text('Correctness + output constraints',8.7,2.9,4.1,.5,16,False,'65717E',PP_ALIGN.CENTER)
# Reward is assigned to the whole generation, including preceding reasoning.
sh=s.shapes.add_shape(MSO_SHAPE.LEFT_ARROW,Inches(2.1),Inches(3.55),Inches(8.6),Inches(.3))
sh.fill.solid();sh.fill.fore_color.rgb=RGBColor.from_string(teal);sh.line.fill.background()
text('RL advantage applied to CoT + response tokens',2.1,3.9,8.6,.5,19,False,teal,PP_ALIGN.CENTER)
for i,b in enumerate(bullets):text('• '+b,.65,4.65+i*1.0,12.0,.93,23)
text('Mechanism: Drori et al., Output Supervision Can Obfuscate the Chain of Thought',.65,7.03,12,.25,11,False,'65717E')
s.notes_slide.notes_text_frame.text=notes+'\n\nSource: '+source+'\n\n'+audit
prs.save(O/'feedback_spillover_slide.pptx')
(O/'feedback_spillover_slide.txt').write_text(title+'\n\n'+'\n'.join('• '+b for b in bullets)+'\n\nPresentation notes\n'+notes+'\n\nMechanism reference\nDrori et al., Output Supervision Can Obfuscate the Chain of Thought\n'+source+'\n\nEvidence details\n'+audit+'\n')
(O/'feedback_spillover_slide.html').write_text('<!doctype html><html><head><meta charset="utf-8"><title>Feedback spillover</title><style>body{font:22px/1.45 Arial;color:#202a36;max-width:1000px;margin:50px auto;padding:20px}h1{font-size:36px}li{margin-bottom:22px}.diagram{padding:28px;background:#e8f1f5;text-align:center}.notes{font-size:19px}small{font-size:15px}</style></head><body><h1>'+html.escape(title)+'</h1><div class="diagram">CoT → Final response → Reward<br><small>Reasoning enabled; reward grades correctness and output constraints</small><br><span style="color:#008f85">← RL advantage applied to CoT + response tokens</span></div><ul>'+''.join('<li>'+html.escape(b)+'</li>' for b in bullets)+'</ul><p class="notes"><b>Presentation notes</b><br>'+html.escape(notes)+'</p><small>Mechanism: <a href="'+source+'">Drori et al., Output Supervision Can Obfuscate the Chain of Thought</a></small></body></html>')
shutil.copy(__file__,O/'build_spillover_slide.py')
# Reopen to verify deliverable count, notes, and object bounds.
r=Presentation(O/'feedback_spillover_slide.pptx');assert len(r.slides)==1
for sh in r.slides[0].shapes:
 assert sh.left>=0 and sh.top>=0 and sh.left+sh.width<=r.slide_width and sh.top+sh.height<=r.slide_height
assert notes in r.slides[0].notes_slide.notes_text_frame.text
assert len(bullets)==2
print('Created and validated one standalone slide, two bullets, five-sentence notes, and review/text copies.')
