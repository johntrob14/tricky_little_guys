from pathlib import Path
import json,html,shutil
import matplotlib
matplotlib.use('Agg')
import matplotlib.pyplot as plt
import numpy as np
from pptx import Presentation
from pptx.util import Inches,Pt
from pptx.dml.color import RGBColor
O=Path('outputs');d=json.loads((O/'output_transfer_data.json').read_text());m=d['models']
title='Training CoT Control Teaches Output Control'
bullets=[
 'These models were trained to control their CoT; we now test instruction following in the visible response with thinking disabled.',
 'Transfer across channels suggests that CoT training improves instruction-following skills shared with the visible response.'
]
notes=(
 'We reuse GPT-OSS-120B trained with GRPO on six CoT rules, the requested reasoning language, or a reasoning token budget. '
 'At evaluation, we move the rule instructions onto the visible answer and explanation and force the analysis channel empty. '
 'Across the same 480 question–rule items, strict output compliance is 44.3% for base, 80.6% after six-rule training, 49.2% after language training, and 59.3% after token-budget training. '
 'The largest gain comes from six-rule training, supporting the interpretation that some instruction-following skills transfer across channels. '
 'These are surface-compliance results from single training runs: RL updated the full generated sequence and included final-answer feedback, while shorter responses may contribute to the token-budget result.'
)
caption=(
 'GPT-OSS-120B; eight deterministic rule types, 480 items, equally weighted after averaging samples per item. '
 'Strict grading covers the entire visible response, allows meta discussion, and does not require a correct answer. '
 '95% confidence intervals: paired bootstrap over 413 normalized-question clusters, 5,000 draws, seed 20261030; evaluation uncertainty only. '
 'Base, grpo3 s100, and len2 s60: 1,632 generations each; lang1 s40: 960. '
 'Language checkpoint is step 40 on both this slide and the preceding CoT-control slide. '
 'The output prompt says “State your answer, and then explain it.” The answer-format restriction is removed and the constraint is scoped to the response. '
 'Code uses temperature 1.0, top-p 1.0, an 8k token cap, and an empty-analysis/ open-final prefill. '
 'This establishes within-model output-control changes, not a channel-only causal isolation or preserved answer quality.'
)
plt.rcParams.update({'font.family':'DejaVu Sans','font.size':19,'svg.fonttype':'none','axes.spines.top':False,'axes.spines.right':False,'axes.spines.left':False,'text.color':'#202A36','axes.labelcolor':'#202A36','xtick.color':'#202A36','ytick.color':'#202A36'})
fig,ax=plt.subplots(figsize=(16,4));fig.subplots_adjust(left=.10,right=.985,bottom=.24,top=.97)
order=['base','grpo3','lang1-s40','len2'];v=np.array([m[x]['rate']*100 for x in order]);ci=np.array([m[x]['ci95'] for x in order])*100
ax.bar(range(4),v,width=.57,color=['#8995A3','#3467A4','#008F85','#D67B35'],zorder=3)
ax.errorbar(range(4),v,yerr=np.array([v-ci[:,0],ci[:,1]-v]),fmt='none',ecolor='#202A36',elinewidth=1.8,capsize=6,zorder=4)
for i in range(4):ax.text(i,ci[i,1]+2,f'{v[i]:.1f}%',ha='center',va='bottom',fontsize=21)
ax.set_xticks(range(4),['Base','Six CoT rules\nGRPO','Reasoning language\nGRPO · step 40','CoT token budget\nGRPO']);ax.set_ylabel('Strict Compliance');ax.set_ylim(0,100);ax.set_yticks([0,20,40,60,80,100],['0%','20%','40%','60%','80%','100%']);ax.grid(axis='y',color='#E5E8EC',zorder=0);ax.tick_params(length=0,pad=10)
fig.savefig(O/'cot_to_output_control.svg',transparent=True);fig.savefig(O/'cot_to_output_control.png',dpi=160,facecolor='white');plt.close(fig)
prs=Presentation();prs.slide_width=Inches(13.333);prs.slide_height=Inches(7.5);s=prs.slides.add_slide(prs.slide_layouts[6])
def text(t,x,y,w,h,size=23,bold=False,color='202A36'):
 sh=s.shapes.add_textbox(Inches(x),Inches(y),Inches(w),Inches(h));tf=sh.text_frame;tf.word_wrap=True;tf.margin_left=tf.margin_right=Inches(.02)
 p=tf.paragraphs[0];p.text=t;p.font.name='Aptos';p.font.size=Pt(size);p.font.bold=bold;p.font.color.rgb=RGBColor.from_string(color)
text(title,.65,.35,12,1.05,32,True)
text('GPT-OSS-120B · 480 items · reasoning disabled at evaluation · 95% CIs',.65,1.1,12,.35,15,color='65717E')
s.shapes.add_picture(str(O/'cot_to_output_control.png'),Inches(.65),Inches(1.55),width=Inches(12))
for j,t in enumerate(bullets):text('• '+t,.65,4.85+j*.85,12,.78)
text('Same checkpoints as preceding slide: six-rule GRPO s100, language GRPO s40, token-budget GRPO s60.',.65,6.95,12,.3,12,color='65717E')
s.notes_slide.notes_text_frame.text=notes+'\n\n'+caption+'\n\nPaired gains versus base (percentage points, 95% CI):\n'+'\n'.join(f"{x}: +{m[x]['delta']*100:.1f} [{m[x]['delta_ci95'][0]*100:.1f}, {m[x]['delta_ci95'][1]*100:.1f}]" for x in order[1:])+'\n\nAudit: output_transfer_data.json, output_transfer_verification.txt, output_transfer_recount.py.'
prs.save(O/'training_cot_control_teaches_output_control.pptx')
(O/'training_cot_control_teaches_output_control.txt').write_text(title+'\n\n'+'\n'.join('• '+t for t in bullets)+'\n\nDescription / speaker notes\n'+notes+'\n\nFigure caption\n'+caption+'\n')
(O/'training_cot_control_teaches_output_control.html').write_text('<!doctype html><html><head><meta charset="utf-8"><title>'+title+'</title><style>body{font:22px/1.4 Arial;color:#202a36;max-width:1120px;margin:35px auto;padding:25px}h1{font-size:36px}svg{width:100%;height:auto}li{margin:16px 0}p{font-size:19px}small{font-size:14px}</style></head><body><h1>'+title+'</h1><small>GPT-OSS-120B · 480 items · reasoning disabled at evaluation · 95% CIs</small>'+(O/'cot_to_output_control.svg').read_text()+'<ul>'+''.join('<li>'+html.escape(t)+'</li>' for t in bullets)+'</ul><p><b>Description / speaker notes</b><br>'+html.escape(notes)+'</p><small>'+html.escape(caption)+'</small></body></html>')
for f in ['output_transfer_recount.py','build_output_transfer.py']:shutil.copy(Path('work')/f,O/f)
r=Presentation(O/'training_cot_control_teaches_output_control.pptx');assert len(r.slides)==1
for sh in r.slides[0].shapes:assert sh.left>=0 and sh.top>=0 and sh.left+sh.width<=r.slide_width and sh.top+sh.height<=r.slide_height
print('Created one slide, plot with confidence intervals, description, notes, and audit files.')
